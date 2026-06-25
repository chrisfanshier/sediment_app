import streamlit as st
import os
import pandas as pd
import io
from db_util import get_connection, get_mgr_export_data
from datetime import date
from sqlalchemy import text
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
import re


def _write_table_sheet(workbook, sheet_name, headers, rows):
    """Add a simple tabular worksheet with bold headers."""
    ws = workbook.create_sheet(sheet_name)

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)

    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col_idx, header in enumerate(headers, 1):
        max_length = len(str(header))
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=2):
            value = row[0].value
            if value is not None:
                max_length = max(max_length, len(str(value)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_length + 2, 60)

    return ws


def _get_ct_export_rows(cruise_id):
    """Return the CT Manager Database View summary filtered to one cruise."""
    with get_connection() as conn:
        result = conn.execute(text("""
            SELECT DISTINCT
                x.core_name,
                x.section_number,
                x.xml_root_tag,
                x.image_type,
                x.processing_stage
            FROM ct_xml x
            LEFT JOIN sections s ON x.section_id = s.section_id
            JOIN cores c ON x.core_id = c.core_id
                OR s.core_id = c.core_id
                OR x.core_name = c.core_name
            WHERE c.cruise_id = :cruise_id
            ORDER BY x.core_name, x.section_number
        """), {"cruise_id": cruise_id})
        rows = [dict(row._mapping) for row in result.fetchall()]

    if not rows:
        return []

    df = pd.DataFrame(rows)
    directory_types = ["CT-Volume", "Helical", "Laminography"]
    image_types = ["giXRayImage", "giCTOrthogonalView"]
    summary_rows = []

    all_pairs = df[["core_name", "section_number"]].drop_duplicates()
    for _, row in all_pairs.iterrows():
        core = row["core_name"]
        section = row["section_number"]
        subdf = df[(df["core_name"] == core) & (df["section_number"] == section)]
        xml_count = len(subdf)

        ct_scan_records = subdf[subdf["xml_root_tag"].isin(["geotek-ct-stack", "geotek-ct-recon"])]
        directories = []
        for dtype in directory_types:
            stages = ct_scan_records.loc[
                ct_scan_records["image_type"] == dtype,
                "processing_stage"
            ].dropna().unique()
            if len(stages) > 0:
                directories.append(f"{dtype} ({', '.join(sorted(stages))})")

        image_records = subdf[subdf["xml_root_tag"] == "geotek-image"]
        images = [itype for itype in image_types if (image_records["image_type"] == itype).any()]

        summary_rows.append((
            core,
            section,
            xml_count,
            "; ".join(directories) if directories else "",
            ", ".join(images) if images else "",
        ))

    return sorted(summary_rows, key=lambda item: (str(item[0]), str(item[1])))


def _get_mscl_export_rows(cruise_id):
    """Return one row per core with MSCL file names comma-space concatenated."""
    with get_connection() as conn:
        result = conn.execute(text("""
            SELECT
                c.core_name,
                COALESCE(
                    string_agg(mf.file_name, ', ' ORDER BY mf.file_name)
                    FILTER (WHERE mf.file_id IS NOT NULL),
                    ''
                ) AS file_name
            FROM cores c
            LEFT JOIN mscl_files mf ON mf.core_id = c.core_id
            WHERE c.cruise_id = :cruise_id
            GROUP BY c.core_id, c.core_name
            ORDER BY c.core_name
        """), {"cruise_id": cruise_id})
        return [(row.core_name, row.file_name) for row in result.fetchall()]


def _get_core_media_export_rows(cruise_id):
    """Return one row per core with selected core_media fields concatenated."""
    with get_connection() as conn:
        result = conn.execute(text("""
            SELECT
                c.core_name,
                COALESCE(
                    string_agg(cm.file_name, ', ' ORDER BY cm.file_name)
                    FILTER (WHERE cm.media_id IS NOT NULL),
                    ''
                ) AS file_name,
                COALESCE(
                    string_agg(CAST(cm.file_size_bytes AS TEXT), ', ' ORDER BY cm.file_name)
                    FILTER (WHERE cm.file_size_bytes IS NOT NULL),
                    ''
                ) AS file_size_bytes,
                COALESCE(
                    string_agg(cm.remarks, ', ' ORDER BY cm.file_name)
                    FILTER (WHERE cm.remarks IS NOT NULL AND cm.remarks <> ''),
                    ''
                ) AS remarks,
                COALESCE(
                    string_agg(cm.media_type, ', ' ORDER BY cm.file_name)
                    FILTER (WHERE cm.media_type IS NOT NULL),
                    ''
                ) AS media_type
            FROM cores c
            LEFT JOIN core_media cm ON cm.core_id = c.core_id
            WHERE c.cruise_id = :cruise_id
            GROUP BY c.core_id, c.core_name
            ORDER BY c.core_name
        """), {"cruise_id": cruise_id})
        return [
            (row.core_name, row.file_name, row.file_size_bytes, row.remarks, row.media_type)
            for row in result.fetchall()
        ]


def view_cruises():
    """Main function to display cruise overview page with cruise selection"""
    
    # --- View Cruise Data Page ---
    st.title("🚢 View Cruise Data")

    # Initialize session state
    if "selected_cruise" not in st.session_state:
        st.session_state.selected_cruise = None

    # --- Cruise Selection ---
    st.subheader("Select Cruise to View")

    # Get all cruise names for selection
    with get_connection() as conn:
        cruises_query = """
        SELECT cr.cruise_id, cr.cruise_name, v.vessel_name 
        FROM cruises cr
        LEFT JOIN vessels v ON cr.vessel_id = v.vessel_id
        ORDER BY cr.cruise_name
        """
        all_cruises = conn.execute(text(cruises_query)).fetchall()

    if not all_cruises:
        st.warning("No cruises found in the database.")
        st.stop()

    # Create cruise options with vessel info for display
    cruise_display_options = {}
    cruise_id_map = {}
    for cruise in all_cruises:
        vessel_info = f" ({cruise[2]})" if cruise[2] else ""
        display_name = f"{cruise[1]}{vessel_info}"
        cruise_display_options[display_name] = cruise[0]
        cruise_id_map[cruise[0]] = cruise[1]  # Map ID to cruise name

    selected_cruise_display = st.selectbox("Select Cruise to View", 
                                         ["Select a cruise..."] + list(cruise_display_options.keys()))

    if selected_cruise_display == "Select a cruise...":
        st.info("Please select a cruise to view.")
        st.stop()

    # Load the selected cruise data
    selected_cruise_id = cruise_display_options[selected_cruise_display]

    with get_connection() as conn:
        # Get cruise data with vessel info
        cruise_query = """
            SELECT cr.*, v.vessel_name
            FROM cruises cr
            LEFT JOIN vessels v ON cr.vessel_id = v.vessel_id
            WHERE cr.cruise_id = :cruise_id
        """
        cruise_data = conn.execute(text(cruise_query), {"cruise_id": selected_cruise_id}).fetchone()
        
        if not cruise_data:
            st.error("Cruise not found!")
            st.stop()

    # --- Display Cruise Data (Clean Display) ---
    st.subheader("Cruise Information")
    st.info(f"**{cruise_data.cruise_name}**")

    # Display cruise info in columns without text boxes
    col1, col2 = st.columns(2)
    
    with col1:
        st.markdown("### 🚢 Vessel & Personnel")
        if cruise_data.vessel_name:
            st.write(f"**Vessel:** {cruise_data.vessel_name}")
        if cruise_data.chief_scientist:
            st.write(f"**Chief Scientist:** {cruise_data.chief_scientist}")
        
        st.markdown("### 🏴 Ports")
        if cruise_data.departure_port:
            st.write(f"**Departure:** {cruise_data.departure_port}")
        if cruise_data.arrival_port:
            st.write(f"**Arrival:** {cruise_data.arrival_port}")

    with col2:
        st.markdown("### 📅 Dates")
        if cruise_data.start_date:
            st.write(f"**Start Date:** {cruise_data.start_date.strftime('%Y-%m-%d')}")
        if cruise_data.end_date:
            st.write(f"**End Date:** {cruise_data.end_date.strftime('%Y-%m-%d')}")
            if cruise_data.start_date:
                duration = (cruise_data.end_date - cruise_data.start_date).days
                st.write(f"**Duration:** {duration} days")
        
        st.markdown("### 💰 Funding")
        if cruise_data.funding_agency:
            st.write(f"**Agency:** {cruise_data.funding_agency}")

    # Description and URL
    if cruise_data.cruise_description:
        st.markdown("### 📝 Description")
        st.write(cruise_data.cruise_description)

    if cruise_data.cruise_report_url:
        st.markdown("### 📄 Cruise Report")
        st.write(f"[View Report]({cruise_data.cruise_report_url})")

    # --- Enhanced Related Records Section ---
    if cruise_data:
        st.divider()
        st.subheader("📊 Related Records")
        
        # Show related cores with file information
        with get_connection() as conn:
            # Enhanced query to show file availability
            related_cores_query = """
            SELECT 
                c.core_name, 
                c.date, 
                c.station,
                COUNT(DISTINCT sf.file_id) as sensor_files,
                COUNT(DISTINCT mf.file_id) as mscl_files,
                CASE WHEN COUNT(DISTINCT sf.file_id) > 0 THEN '✅' ELSE '❌' END as has_sensor,
                CASE WHEN COUNT(DISTINCT mf.file_id) > 0 THEN '✅' ELSE '❌' END as has_mscl,
                CASE
                    WHEN EXISTS (
                        SELECT 1
                        FROM ct_xml x
                        LEFT JOIN sections s ON x.section_id = s.section_id
                        WHERE x.core_id = c.core_id OR s.core_id = c.core_id
                    ) THEN '✅' ELSE '❌'
                END as has_ct
            FROM cores c
            LEFT JOIN sensor_files sf ON c.core_id = sf.core_id
            LEFT JOIN mscl_files mf ON c.core_id = mf.core_id
            WHERE c.cruise_id = :cruise_id 
            GROUP BY c.core_id, c.core_name, c.date, c.station
            ORDER BY c.core_name
            """
            related_cores = conn.execute(text(related_cores_query), {"cruise_id": selected_cruise_id}).fetchall()
            
            # Get related winch data count
            winch_count_query = "SELECT COUNT(*) FROM winch_data WHERE cruise_id = :cruise_id"
            winch_count = conn.execute(text(winch_count_query), {"cruise_id": selected_cruise_id}).scalar()
        
        if related_cores:
            st.write(f"**Cores from this cruise ({len(related_cores)}):**")
            
            # Create enhanced dataframe
            cores_df = pd.DataFrame(related_cores, columns=[
                'Core Name', 'Date', 'Station', 'Sensor Files', 'MSCL Files', 'Sensor Data', 'MSCL Data', 'CT Data'
            ])
            
            # Display the enhanced table
            st.dataframe(
                cores_df,
                use_container_width=True,
                column_config={
                    "Core Name": st.column_config.TextColumn("Core Name", width="medium"),
                    "Date": st.column_config.DateColumn("Date", width="small"),
                    "Station": st.column_config.TextColumn("Station", width="small"),
                    "Sensor Files": st.column_config.NumberColumn("# Sensor Files", width="small"),
                    "MSCL Files": st.column_config.NumberColumn("# MSCL Files", width="small"),
                    "Sensor Data": st.column_config.TextColumn("Sensor", width="small"),
                    "MSCL Data": st.column_config.TextColumn("MSCL", width="small"),
                    "CT Data": st.column_config.TextColumn("CT", width="small"),
                }
            )
            
            # Summary statistics
            total_sensor_files = sum(row[3] for row in related_cores)
            total_mscl_files = sum(row[4] for row in related_cores)
            cores_with_sensor = sum(1 for row in related_cores if row[3] > 0)
            cores_with_mscl = sum(1 for row in related_cores if row[4] > 0)
            cores_with_ct = sum(1 for row in related_cores if row[7] == '✅')
            
            st.markdown("#### 📈 Data Summary")
            summary_col1, summary_col2, summary_col3, summary_col4, summary_col5 = st.columns(5)
            
            with summary_col1:
                st.metric("Total Sensor Files", total_sensor_files)
            with summary_col2:
                st.metric("Total MSCL Files", total_mscl_files)
            with summary_col3:
                st.metric("Cores with Sensor Data", f"{cores_with_sensor}/{len(related_cores)}")
            with summary_col4:
                st.metric("Cores with MSCL Data", f"{cores_with_mscl}/{len(related_cores)}")
            with summary_col5:
                st.metric("Cores with CT Data", f"{cores_with_ct}/{len(related_cores)}")
                
        else:
            st.info("No cores found for this cruise.")
        
        if winch_count > 0:
            st.write(f"**Winch Data Records:** {winch_count}")
        else:
            st.info("No winch data records found for this cruise.")

# --- MGR Excel Export ---
    st.divider()
    st.subheader("📥 Export MGR Metadata")

    if st.button("Generate OSU-MGR Excel Export"):
        rows = get_mgr_export_data(cruise_data.cruise_name)

        if not rows:
            st.warning("No section data found for this cruise. Add sections before exporting.")
        else:
            wb = Workbook()

            # --- Sheet 2: Dropdown Menus (must exist before sheet1 DVs reference it) ---
            ws2 = wb.create_sheet("Dropdown Menus")
            core_types = [
                ("Box Core", "BC"), ("Gravity Core", "GC"), ("Jumbo Piston Core", "JC"),
                ("Kasten Core", "KC"), ("Multicore", "MC"), ("Piston Core", "PC"),
                ("Trigger Core", "TC"), ("Vibracore", "VC"),
            ]
            working_archive = ["Working", "Archive", "Round"]
            for i, (name, abbr) in enumerate(core_types, 1):
                ws2.cell(row=i, column=1, value=name)
                ws2.cell(row=i, column=2, value=abbr)
            for i, val in enumerate(working_archive, 1):
                ws2.cell(row=i, column=3, value=val)

            # --- Sheet 1: For PI ---
            ws1 = wb.active
            ws1.title = "For PI"

            # Header block rows 1-5
            first = rows[0]
            ws1["A1"], ws1["B1"] = "Cruise Name", first[0]
            ws1["A2"], ws1["B2"] = "RV Name", first[1]
            ws1["A3"], ws1["B3"] = "PI Scientist", first[2]
            ws1["A5"] = None  # spacer row

            # Column headers row 6
            headers = [
                "Core Number", "Recovery Method", "Core Abbreviation",
                "Date Collected", "Time collected (UTC)", "Decimal Latitude",
                "Decimal Longitude", "Water Depth (m)", "Section Number",
                "Working/Archive", "Section Name", "Alternate Section Name or Ship Station",
                "Section Length (cm)", "Top of section interval (cm)",
                "Bottom of section interval (cm)", "Core Diameter (cm)",
                "Core Tube Material", "Notes"
            ]
            for col, header in enumerate(headers, 1):
                cell = ws1.cell(row=6, column=col, value=header)
                cell.font = Font(bold=True)

            # Data rows starting at row 7 (columns 4 onward in query = index 3+)
            for row_idx, row in enumerate(rows, 7):
                data = [
                    row[3],   # core_name          A
                    row[4],   # recovery_method     B
                    row[5],   # core_abbreviation   C
                    row[6],   # date_collected      D
                    row[7],   # time_collected_utc  E
                    row[8],   # latitude            F
                    row[9],   # longitude           G
                    row[10],  # water_depth_m       H
                    row[11],  # section_number      I
                    row[12],  # working_archive     J
                    row[13],  # section_name        K
                    row[14],  # alternate_name      L
                    row[15],  # section_length      M
                    row[16],  # top_cum_length      N
                    row[17],  # bottom_cum_length   O
                    row[18],  # core_diameter       P
                    row[19],  # core_tube_material  Q
                    row[20],  # remarks             R
                ]
                for col_idx, val in enumerate(data, 1):
                    ws1.cell(row=row_idx, column=col_idx, value=val)

            # Data validation dropdowns
            dv_recovery = DataValidation(
                type="list",
                formula1="'Dropdown Menus'!$A$1:$A$8",
                allow_blank=True, showInputMessage=True, showErrorMessage=True
            )
            dv_working_archive = DataValidation(
                type="list",
                formula1="'Dropdown Menus'!$C$1:$C$3",
                allow_blank=True, showInputMessage=True, showErrorMessage=True
            )
            ws1.add_data_validation(dv_recovery)
            ws1.add_data_validation(dv_working_archive)
            dv_recovery.sqref = "B7:B1048576"
            dv_working_archive.sqref = "J7:J1048576"

            _write_table_sheet(
                wb,
                "CT_Data",
                ["Core", "Section", "Total XMLs", "CT/Scan Types", "Image Types"],
                _get_ct_export_rows(selected_cruise_id)
            )
            _write_table_sheet(
                wb,
                "MSCL_Data",
                ["core_name", "file_name"],
                _get_mscl_export_rows(selected_cruise_id)
            )
            _write_table_sheet(
                wb,
                "core_media",
                ["core_name", "file_name", "file_size_bytes", "remarks", "media_type"],
                _get_core_media_export_rows(selected_cruise_id)
            )

            # Serialize to bytes and offer download
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)

            st.download_button(
                label="⬇️ Download Excel",
                data=buf,
                file_name=f"OSU-MGR_Marine_Sediment_Data-{cruise_data.cruise_name}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )