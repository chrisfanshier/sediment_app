#!/usr/bin/env python3
"""Build docs/sensorius_math_reference.docx with Word-native equations."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt

OUT = Path(__file__).parent / "sensorius_math_reference.docx"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _shade(cell, fill: str) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(el)


def h(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def p(doc: Document, text: str, *, bold: bool = False, italic: bool = False) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    para.paragraph_format.space_after = Pt(4)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def eq(doc: Document, inner: str) -> None:
    """Insert centered Word OMML equation (inner = m:* elements only)."""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(4)
    omath = parse_xml(f"<m:oMath {nsdecls('m')}>{inner}</m:oMath>")
    para._p.append(omath)


def tbl(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, head in enumerate(headers):
        c = table.rows[0].cells[i]
        c.text = head
        for r in c.paragraphs[0].runs:
            r.bold = True
        _shade(c, "D9E2F3")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = val
    doc.add_paragraph()


def frac(num: str, den: str) -> str:
    return f"<m:f><m:num><m:r><m:t>{num}</m:t></m:r></m:num><m:den><m:r><m:t>{den}</m:t></m:r></m:den></m:f>"


def sub(base: str, subscript: str) -> str:
    return (
        f"<m:sSub><m:e><m:r><m:t>{base}</m:t></m:r></m:e>"
        f"<m:sub><m:r><m:t>{subscript}</m:t></m:r></m:sub></m:sSub>"
    )


def sup(base: str, superscript: str) -> str:
    return (
        f"<m:sSup><m:e><m:r><m:t>{base}</m:t></m:r></m:e>"
        f"<m:sup><m:r><m:t>{superscript}</m:t></m:r></m:sup></m:sSup>"
    )


def txt(*parts: str) -> str:
    return "".join(f"<m:r><m:t>{x}</m:t></m:r>" for x in parts)


def build() -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_heading("Sensorius — Mathematical Reference", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p(
        doc,
        "Units, algorithms, and output definitions for Sensorius. Each analysis mode "
        "includes a summary, step-by-step procedure, and formulas as implemented in "
        "sensor_standalone. Persistence describes the portable SQLite project model.",
    )

    # ── 1 Units ───────────────────────────────────────────────────────
    h(doc, "1. Units and conventions")
    tbl(
        doc,
        ["Quantity", "Storage / UI", "In calculations"],
        [
            ["Depths", "metres (m)", "metres"],
            ["Scope, core length, trigger core length", "feet (ft)", "Lₘ = Lft ÷ 3.28"],
            ["Trigger line length (Alter)", "feet (ft)", "Lₘ = Lft ÷ 3.28"],
            ["Trigger penetration", "metres (m)", "metres"],
            ["Time offsets", "seconds (s)", "seconds"],
            ["Sampling rate", "hertz (Hz)", "fₛ = 1 / median(Δt)"],
            ["Weight-stand length", "—", "1.5 m (fixed)"],
            ["Piston offset constant", "metres (m)", "default 1.25 m"],
        ],
    )
    bullets(
        doc,
        [
            "Depth increases downward (oceanographic convention).",
            "FT_TO_M = 1/3.28 throughout (not 0.3048).",
            "Indices refer to merged sensor dataframe rows.",
            "NaNs are interpolated before filtering or geometry.",
        ],
    )

    # ── 2 Preprocessing ───────────────────────────────────────────────
    h(doc, "2. Preprocessing (Data Manager)")
    p(doc, "Summary", bold=True)
    p(
        doc,
        "Parse RBR/Star-Oddi files, keep channels matching depth/pressure/acceleration, "
        "prefix columns by sensor location and filename, outer-merge on datetime.",
    )
    p(doc, "Algorithm", bold=True)
    numbered(
        doc,
        [
            "Parse each file → datetime + numeric columns.",
            "Rename: {location}_{filename}_{column}.",
            "Outer merge on datetime; sort ascending.",
            "Depth columns contain 'Depth' in the name.",
        ],
    )
    p(doc, "Winch auto-load", bold=True)
    numbered(
        doc,
        [
            "Select winch files overlapping core deployment window.",
            "Concatenate, sort by datetime, trim to deployment window.",
            "Manual winch mode loads full files without trimming.",
        ],
    )

    # ── 3 Modes ───────────────────────────────────────────────────────
    h(doc, "3. Analysis modes")

    h(doc, "3.1 View Data", 2)
    p(doc, "Summary: Display merged depth time series.")
    p(doc, "Outputs: Visual only.")

    h(doc, "3.2 Depth Offset", 2)
    p(doc, "Summary: Depth-dependent calibration regression and/or manual offsets.")
    p(doc, "Regression", bold=True)
    eq(doc, sub("Δ", "j,i") + txt("(d) = slope · d + intercept"))
    p(doc, "Δj,i = predicted (sensorⱼ − sensorᵢ) at depth d.")
    eq(
        doc,
        txt("d")
        + sub("", "target,corr")
        + txt(" = d")
        + sub("", "target")
        + txt(" + sign · ")
        + sub("Δ", "j,i")
        + txt("(d")
        + sub("", "ref")
        + txt(")"),
    )
    tbl(
        doc,
        ["Reference", "Sign", "Action"],
        [
            ["sensorᵢ in pair (j−i)", "+1", "Add offset to target i"],
            ["sensorⱼ in pair (j−i)", "−1", "Subtract offset from target j"],
        ],
    )
    p(doc, "Manual offset", bold=True)
    eq(doc, txt("d") + sub("", "col,corr") + txt(" = d") + sub("", "col") + txt(" + ") + sub("δ", "m"))
    p(doc, "δₘ in metres.")

    h(doc, "3.3 Time Offset", 2)
    p(doc, "Summary: Heave cross-correlation for clock alignment.")
    p(doc, "Heave: 2nd-order detrend, then 4th-order Butterworth bandpass 0.05–0.5 Hz.", bold=True)
    eq(
        doc,
        txt("c[k] = Σ")
        + sub("", "n")
        + sub("h", "ref")
        + txt("[n] · ")
        + sub("h", "mov")
        + txt("[n−k]"),
    )
    eq(doc, txt("k* = argmax |c[k]|"))
    eq(doc, txt("δ = ½ · ") + frac("c[k*−1] − c[k*+1]", "c[k*−1] − 2c[k*] + c[k*+1]"))
    eq(doc, txt("Δt = (k* + δ) / ") + sub("f", "s"))
    p(doc, "Apply via linear interpolation: z_corrected(t) = interp(t, t+Δt, z).")

    h(doc, "3.4 Create Calibration", 2)
    p(doc, "Summary: ≥2 casts → linear regression of pairwise differences vs mean depth.")
    eq(doc, sub("d̄", "all") + txt(" = meanₙ(meanᵢ dᵢ[n])"))
    eq(doc, sub("Δ̄", "j,i") + txt(" = meanₙ(dⱼ[n] − dᵢ[n])"))
    eq(doc, sub("Δ̄", "j,i") + txt(" = slope · ") + sub("d̄", "all") + txt(" + intercept"))

    h(doc, "3.5 Trip Detector", 2)
    p(doc, "Summary: Trip when cross-sensor derivative spread exceeds threshold.")
    eq(
        doc,
        sup(sub("d̂", "i"), "(k)")
        + txt("[n] = Savgol(dᵢ, W, P, deriv=k, Δt)"),
    )
    p(doc, "Defaults: W=51, P=3, k=1 (velocity), fₛ=32 Hz, τ=0.5, edge buffer=500.")
    eq(doc, txt("σ[n] = stdᵢ( ") + sup(sub("d̂", "i"), "(k)") + txt("[n] )"))
    numbered(
        doc,
        [
            "Mask first/last 500 samples.",
            "If σ[n] > τ: trip_idx = first n (confidence 0.8).",
            "Else: trip_idx = argmax σ (confidence 0.5).",
        ],
    )

    h(doc, "3.6 Piston Position", 2)
    p(doc, "Start-core (Sₘ = scopeft / 3.28):", bold=True)
    eq(
        doc,
        txt("start_core = min{n ≥ trip : |")
        + sub("d", "release")
        + txt("[n] − ")
        + sub("d", "WS")
        + txt("[n]| > S")
        + sub("", "m")
        + txt("}"),
    )
    p(doc, "Piston trace (Bₘ = core_lengthft / 3.28, O = 1.25 m default):", bold=True)
    eq(doc, sub("d", "piston") + txt("[n] = ") + sub("d", "WS") + txt("[n] + O + B") + sub("", "m"))
    eq(
        doc,
        sub("d", "piston")
        + txt("[n] = ")
        + sub("d", "release")
        + txt("[n] + S")
        + sub("", "m")
        + txt(" + B")
        + sub("", "m")
        + txt(" + O   (from start_core onward)"),
    )

    h(doc, "3.7 Winch", 2)
    p(doc, "Trip: first w[n−1] − w[n] ≥ 1000 in search window (edge buffer 500).")
    eq(doc, sub("Δt", "winch") + txt(" = t") + sub("", "sensor") + txt(" − t") + sub("", "winch,orig"))
    eq(
        doc,
        txt("T")
        + sub("", "winch,aligned")
        + txt("[n] = T")
        + sub("", "winch,orig")
        + txt("[n] + ")
        + sub("Δt", "winch"),
    )

    h(doc, "3.8 Calculate", 2)
    p(doc, "Seafloor:", bold=True)
    eq(
        doc,
        sub("d", "seafloor")
        + txt(" = ")
        + sub("d", "trigger")
        + txt("[trip] + L")
        + sub("", "TC,m")
        + txt(" − ")
        + sub("p", "trigger"),
    )
    p(doc, "Core tip & start_pen:", bold=True)
    eq(doc, sub("d", "tip") + txt("[n] = ") + sub("d", "WS") + txt("[n] + 1.5 + L") + sub("", "core,m"))
    eq(
        doc,
        txt("start_pen = min{n ≥ trip : ")
        + sub("d", "tip")
        + txt("[n] ≥ ")
        + sub("d", "seafloor")
        + txt("}"),
    )
    tbl(
        doc,
        ["Output", "Formula", "Requires"],
        [
            ["recoil_max", "|d_rel[trip] − min d_rel[trip…end_pen]|", "trip"],
            ["recoil_time_s", "(n_peak − trip) × median(Δt)", "trip"],
            ["fall_dist", "|d_WS[trip] − d_WS[end_pen]|", "trip"],
            ["suck_in", "|d_WS[end_pen] − max d_WS|", "trip"],
            ["recoil_start", "|d_rel[trip] − d_rel[start_core]|", "trip + start_core"],
            ["freefall_start", "|d_WS[trip] − d_WS[start_core]|", "trip + start_core"],
            ["piston_suck", "d_piston[pullout] − d_piston[cross]", "piston"],
            ["pen_deficit", "d_seafloor − max d_WS", "seafloor"],
            ["piston_alt", "d_seafloor − d_piston[start_core]", "seafloor"],
            ["freefall_est", "|d_WS[trip]+1.5+L_core,m−d_seafloor|", "seafloor"],
        ],
    )

    h(doc, "3.9 Alter", 2)
    p(doc, "Summary: Re-run piston + Calculate with edited rig parameters; trip/end_pen/pullout fixed.")
    eq(doc, sub("Δ", "trig,m") + txt(" = (trigger_line_ft_new − trigger_line_ft_db) / 3.28"))
    eq(doc, sub("Δ", "TC,m") + txt(" = (trigger_core_length_ft_new − trigger_core_length_ft_db) / 3.28"))
    eq(
        doc,
        txt("d")
        + sub("", "WS,mod")
        + txt(" = d")
        + sub("", "WS")
        + txt(" − ")
        + sub("Δ", "trig,m")
        + txt(" − ")
        + sub("Δ", "TC,m"),
    )
    p(doc, "Similarly: d_release,mod and d_trigger,mod with the same shifts.")
    p(doc, "Pig-weight changes: not implemented.")

    h(doc, "3.10 Piston Analysis", 2)
    eq(
        doc,
        sub("x", "core")
        + txt("[n] = ")
        + sub("d", "WS")
        + txt("[n] − ")
        + sub("d", "WS")
        + txt("[start_pen]"),
    )
    eq(doc, txt("v[n] = Savgol(d, W, P, deriv=1, Δt)   [m/s],  W=51, P=3"))

    # ── 4 Persistence ─────────────────────────────────────────────────
    h(doc, "4. Save, restore, and persistence")
    p(
        doc,
        "Project folder: sensorius_data.sqlite + sensor_data/<cruise>/ + winch_data/<cruise>/. "
        "Multiple analysis_results rows per core.",
    )
    p(doc, "analysis_results flat columns", bold=True)
    tbl(
        doc,
        ["Column", "Meaning"],
        [
            ["trip_index, trip_datetime", "Sensor trip event"],
            ["start_core_index", "Start of coring"],
            ["result_*", "Calculate outputs (m; recoil_time_s in s)"],
            ["trigger_pen_m", "Trigger penetration input (m)"],
            ["winch_trip_index, winch_time_offset_sec", "Winch alignment"],
            ["params", "JSON: depth_offset, time_offset, trip, piston, calc_inputs, winch, …"],
        ],
    )
    p(doc, "Restore order", bold=True)
    numbered(
        doc,
        [
            "Depth corrections → Time corrections → Trip → Winch → Piston → Calculate",
            "Skipped sections reported in status bar.",
            "trigger_core_penetration written back to core row on save.",
        ],
    )

    h(doc, "5. Savitzky–Golay filter")
    eq(doc, txt("x̂[n] = Σ ") + sub("c", "j") + txt(" x[n+j]"))
    p(doc, "Local polynomial least-squares (SciPy savgol_filter).")

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
